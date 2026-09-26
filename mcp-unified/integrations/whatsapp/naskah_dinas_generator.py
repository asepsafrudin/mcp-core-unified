"""
naskah_dinas_generator.py — Engine Generator Naskah Dinas Resmi Ditjen Bina Bangda Kemendagri.
Mendukung generasi otomatis berkas .docx berbasis master template resmi Permendagri No. 1 Tahun 2023.
Menjaga 100% keaslian Kop Surat Vektor, Font Arial 11pt, Margin, dan Spasi Resmi.
"""

import os
import re
import copy
import logging
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, Optional, List

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

logger = logging.getLogger("naskah_dinas_generator")

BASE_TEMPLATES_DIR = Path("/home/aseps/MCP/storage/templates/naskah_dinas")
OUTPUT_DIR = Path("/home/aseps/MCP/storage/naskah_dinas_output")

MONTH_NAMES_ID = [
    "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember"
]

def format_date_id(dt: Optional[datetime] = None) -> str:
    """Format tanggal Indonesia, contoh: '7 September 2026'."""
    if not dt:
        dt = datetime.now()
    return f"{dt.day} {MONTH_NAMES_ID[dt.month - 1]} {dt.year}"

def set_cell_value(cell, text: str, bold: bool = False, italic: bool = False):
    """Menetapkan teks sel tabel dengan font standar Arial 11pt tanpa merusak tata letak."""
    if not cell.paragraphs:
        p = cell.add_paragraph()
    else:
        p = cell.paragraphs[0]
        p.text = ""
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.bold = bold
    run.italic = italic

def replace_text_in_paragraph(p, search_str: str, replace_str: str):
    """Mengganti kemunculan search_str dalam paragraf sambil menjaga run style jika memungkinkan."""
    if search_str in p.text:
        full_text = p.text.replace(search_str, replace_str)
        # Ambil font run pertama jika ada
        font_name = "Arial"
        font_size = Pt(11)
        bold = False
        italic = False
        if p.runs:
            font_name = p.runs[0].font.name or "Arial"
            font_size = p.runs[0].font.size or Pt(11)
            bold = p.runs[0].bold
            italic = p.runs[0].italic
        p.text = full_text
        if p.runs:
            p.runs[0].font.name = font_name
            p.runs[0].font.size = font_size
            p.runs[0].bold = bold
            p.runs[0].italic = italic

class NaskahDinasGenerator:
    """Engine pembuat naskah dinas Word (.docx) berbasis template resmi Ditjen Bina Bangda."""

    def __init__(self, templates_dir: Optional[Path] = None, output_dir: Optional[Path] = None):
        self.templates_dir = templates_dir or BASE_TEMPLATES_DIR
        self.output_dir = output_dir or OUTPUT_DIR
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate_nd_laporan(self, data: Dict[str, Any], output_filename: Optional[str] = None) -> str:
        """
        Menghasilkan Nota Dinas Laporan Hasil Kegiatan / Rapat / FGD.
        Menggunakan template `nota_dinas/template_nd_laporan_kegiatan.docx`.
        """
        tmpl_path = self.templates_dir / "nota_dinas" / "template_nd_laporan_kegiatan.docx"
        if not tmpl_path.exists():
            tmpl_path = self.templates_dir / "nota_dinas" / "master_nd_laporan_kegiatan.docx"
            
        doc = docx.Document(str(tmpl_path))

        # 1. Update Tabel Header (Tabel 0: 8 rows x 3 cols)
        if doc.tables:
            tbl = doc.tables[0]
            set_cell_value(tbl.rows[0].cells[2], data.get("yth", "Bapak Sekretaris Ditjen Bina Pembangunan Daerah"))
            set_cell_value(tbl.rows[1].cells[2], data.get("dari", "Analis Hukum Ahli Madya"))
            set_cell_value(tbl.rows[2].cells[2], data.get("tembusan", "1. Bapak Dirjen Bina Pembangunan Daerah (sebagai laporan)\n2. Direktur Terkait"))
            set_cell_value(tbl.rows[3].cells[2], data.get("tanggal_nd", format_date_id()))
            set_cell_value(tbl.rows[4].cells[2], data.get("nomor_nd", "000.1.5/        /PUU"))
            set_cell_value(tbl.rows[5].cells[2], data.get("sifat", "Segera"))
            set_cell_value(tbl.rows[6].cells[2], data.get("lampiran", "-"))
            set_cell_value(tbl.rows[7].cells[2], data.get("hal", "Laporan Hasil Pelaksanaan Rapat Koordinasi"), bold=True)

        # 2. Update Paragraf Utama & Batang Tubuh
        # Kita sesuaikan paragraf pembuka dan penutup
        clean_hal = data.get("hal", "Kegiatan Rapat Koordinasi")
        for p in doc.paragraphs:
            if "Menindaklanjuti" in p.text and len(p.text) > 20:
                p_text = data.get("paragraf_pembuka")
                if not p_text:
                    p_text = f"Menindaklanjuti pelaksanaan kegiatan koordinasi mengenai {clean_hal}, dengan ini kami laporkan kepada Bapak Sekretaris Ditjen beberapa hal sebagai berikut:"
                p.text = p_text
                if p.runs:
                    p.runs[0].font.name = "Arial"
                    p.runs[0].font.size = Pt(11)

            # Update penandatangan jika disediakan
            if "Lady Diana Handayani" in p.text and data.get("nama_penandatangan"):
                p.text = data.get("nama_penandatangan")
                if p.runs:
                    p.runs[0].font.name = "Arial"
                    p.runs[0].font.size = Pt(11)
                    p.runs[0].bold = True
            if "NIP 19830306" in p.text and data.get("nip"):
                p.text = f"NIP {data.get('nip')}"
                if p.runs:
                    p.runs[0].font.name = "Arial"
                    p.runs[0].font.size = Pt(11)

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        fname = output_filename or f"ND_Laporan_{ts}.docx"
        out_file = self.output_dir / fname
        doc.save(str(out_file))
        logger.info(f"Nota Dinas Laporan successfully generated at {out_file}")
        return str(out_file)

    def generate_nd_pengantar(self, data: Dict[str, Any], output_filename: Optional[str] = None) -> str:
        """
        Menghasilkan Nota Dinas Pengantar dari Eselon III/Fungsional Madya ke Sesditjen Bangda.
        Cocok untuk penyampaian telaahan regulasi / permohonan paraf konsep.
        """
        tmpl_path = self.templates_dir / "nota_dinas" / "template_nd_pengantar_ke_sesditjen.docx"
        if not tmpl_path.exists():
            tmpl_path = self.templates_dir / "nota_dinas" / "master_nd_pengantar_ke_sesditjen.docx"

        doc = docx.Document(str(tmpl_path))

        # Tabel 0 (Header ND Pengantar)
        if doc.tables:
            tbl = doc.tables[0]
            set_cell_value(tbl.rows[0].cells[2], data.get("yth", "Bapak Sekretaris Ditjen Bina Pembangunan Daerah"))
            set_cell_value(tbl.rows[1].cells[2], data.get("dari", "Analis Hukum Ahli Madya"))
            set_cell_value(tbl.rows[2].cells[2], data.get("tembusan", "-"))
            set_cell_value(tbl.rows[3].cells[2], data.get("tanggal_nd", format_date_id()))
            set_cell_value(tbl.rows[4].cells[2], data.get("nomor_nd", "100.4.4.1/       /PUU"))
            set_cell_value(tbl.rows[5].cells[2], data.get("sifat", "Segera"))
            set_cell_value(tbl.rows[6].cells[2], data.get("lampiran", "Satu berkas"))
            set_cell_value(tbl.rows[7].cells[2], data.get("hal", "Permohonan Paraf Konsep Naskah Dinas"), bold=True)

        # Update rujukan disposisi / pembuka
        for p in doc.paragraphs:
            if "Menindaklanjuti disposisi" in p.text:
                ref_text = data.get("paragraf_pembuka")
                if not ref_text:
                    surat_ref = data.get("nomor_surat_masuk", "-")
                    pengirim = data.get("pengirim_surat", "kementerian/lembaga")
                    hal_surat = data.get("hal_surat_masuk", "tersebut")
                    ref_text = (
                        f"Menindaklanjuti disposisi Bapak Sekretaris Ditjen atas Surat {pengirim} Nomor {surat_ref} "
                        f"Hal {hal_surat}, dengan ini kami laporkan kepada Bapak Sekretaris Ditjen beberapa hal sebagai berikut:"
                    )
                p.text = ref_text
                if p.runs:
                    p.runs[0].font.name = "Arial"
                    p.runs[0].font.size = Pt(11)

            if "Substansi Perundang-undangan telah melakukan pencermatan" in p.text:
                if data.get("telaah_puu"):
                    p.text = data.get("telaah_puu")
                    if p.runs:
                        p.runs[0].font.name = "Arial"
                        p.runs[0].font.size = Pt(11)

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        fname = output_filename or f"ND_Pengantar_Sesditjen_{ts}.docx"
        out_file = self.output_dir / fname
        doc.save(str(out_file))
        logger.info(f"Nota Dinas Pengantar successfully generated at {out_file}")
        return str(out_file)

    def generate_nd_sesditjen_ke_dirjen(self, data: Dict[str, Any], output_filename: Optional[str] = None) -> str:
        """
        Menghasilkan Nota Dinas Tingkat Pimpinan dari Sekretaris Ditjen ke Direktur Jenderal Bina Bangda.
        """
        tmpl_path = self.templates_dir / "nota_dinas" / "template_nd_sesditjen_ke_dirjen.docx"
        if not tmpl_path.exists():
            tmpl_path = self.templates_dir / "nota_dinas" / "master_nd_sesditjen_ke_dirjen.docx"

        doc = docx.Document(str(tmpl_path))

        # Tabel 0
        if doc.tables:
            tbl = doc.tables[0]
            set_cell_value(tbl.rows[0].cells[2], data.get("yth", "Bapak Dirjen Bina Pembangunan Daerah"))
            set_cell_value(tbl.rows[1].cells[2], data.get("dari", "Sekretaris Ditjen Bina Pembangunan Daerah"))
            set_cell_value(tbl.rows[2].cells[2], data.get("tanggal_nd", format_date_id()))
            set_cell_value(tbl.rows[3].cells[2], data.get("nomor_nd", "100.4.2/       /Set/Bangda"))
            set_cell_value(tbl.rows[4].cells[2], data.get("sifat", "Segera"))
            set_cell_value(tbl.rows[5].cells[2], data.get("lampiran", "Satu berkas"))
            set_cell_value(tbl.rows[6].cells[2], data.get("hal", "Permohonan Paraf Konsep Naskah Dinas"), bold=True)

        for p in doc.paragraphs:
            if "Menindaklanjuti Surat" in p.text and len(p.text) > 20:
                p_text = data.get("paragraf_pembuka")
                if not p_text:
                    surat_ref = data.get("nomor_surat_masuk", "-")
                    pengirim = data.get("pengirim_surat", "kementerian/lembaga")
                    hal_surat = data.get("hal_surat_masuk", "tersebut")
                    p_text = (
                        f"Menindaklanjuti Surat {pengirim} Nomor {surat_ref} Hal {hal_surat}, "
                        f"dengan hormat dilaporkan kepada Bapak Dirjen hal-hal sebagai berikut:"
                    )
                p.text = p_text
                if p.runs:
                    p.runs[0].font.name = "Arial"
                    p.runs[0].font.size = Pt(11)

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        fname = output_filename or f"ND_Sesditjen_Dirjen_{ts}.docx"
        out_file = self.output_dir / fname
        doc.save(str(out_file))
        logger.info(f"Nota Dinas Sesditjen ke Dirjen successfully generated at {out_file}")
        return str(out_file)

    def generate_surat_dinas_eksternal(self, data: Dict[str, Any], output_filename: Optional[str] = None) -> str:
        """
        Menghasilkan Surat Dinas Keluar Eksternal kepada Kementerian/Lembaga atau Pemerintah Daerah.
        """
        tmpl_path = self.templates_dir / "surat_dinas" / "template_surat_dinas_eksternal.docx"
        if not tmpl_path.exists():
            tmpl_path = self.templates_dir / "surat_dinas" / "master_surat_dinas_eksternal.docx"

        doc = docx.Document(str(tmpl_path))

        # Update metadata di baris-baris awal
        tgl_str = data.get("tanggal_surat", format_date_id())
        nomor_str = data.get("nomor_surat", "100.4.2/        /Bangda")
        sifat_str = data.get("sifat", "Segera")
        lampiran_str = data.get("lampiran", "Satu Berkas")
        hal_str = data.get("hal", "Penyampaian Naskah Dinas")
        yth_str = data.get("yth", "Kepala Lembaga Terkait")

        for p in doc.paragraphs:
            if "Agustus 2026" in p.text:
                p.text = tgl_str
                if p.runs:
                    p.runs[0].font.name = "Arial"
                    p.runs[0].font.size = Pt(11)
            elif "Nomor" in p.text and ":" in p.text:
                p.text = f"Nomor\t:\t{nomor_str}"
                if p.runs:
                    p.runs[0].font.name = "Arial"
                    p.runs[0].font.size = Pt(11)
            elif "Hal" in p.text and ":" in p.text:
                p.text = f"Hal\t:\t{hal_str}"
                if p.runs:
                    p.runs[0].font.name = "Arial"
                    p.runs[0].font.size = Pt(11)
                    p.runs[0].bold = True
            elif "Yth." in p.text:
                p.text = f"Yth. \t{yth_str}"
                if p.runs:
                    p.runs[0].font.name = "Arial"
                    p.runs[0].font.size = Pt(11)
                    p.runs[0].bold = True

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        fname = output_filename or f"Surat_Dinas_Eksternal_{ts}.docx"
        out_file = self.output_dir / fname
        doc.save(str(out_file))
        logger.info(f"Surat Dinas Eksternal successfully generated at {out_file}")
        return str(out_file)

# Singleton Instance
satria_naskah_generator = NaskahDinasGenerator()
