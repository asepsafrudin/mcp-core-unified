import os
import sys
import copy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Add project root to sys.path
sys.path.insert(0, '/home/aseps/MCP/mcp-unified')

def copy_cell_xml(src_cell, dst_cell):
    """Deep copy paragraphs to preserve formatting."""
    for p in dst_cell.paragraphs:
        p._element.getparent().remove(p._element)
    for p in src_cell.paragraphs:
        dst_cell._tc.append(copy.deepcopy(p._p))

def set_cell_vmerge(cell, merge_type):
    """
    Using a more stable method for vertical merge.
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vMerge = tcPr.xpath('w:vMerge')
    if not vMerge:
        vMerge = OxmlElement('w:vMerge')
        tcPr.append(vMerge)
    else:
        vMerge = vMerge[0]
    
    if merge_type == 'restart':
        vMerge.set(qn('w:val'), 'restart')
    else:
        # For 'continue', we remove the val attribute
        if qn('w:val') in vMerge.attrib:
            del vMerge.attrib[qn('w:val')]

def rebuild_and_merge_fixed(input_path, output_path):
    print(f"Phase 1: Rebuilding content from {input_path}")
    doc_src = Document(input_path)
    table_src = doc_src.tables[0]
    
    agencies = [
        "KEMENKOPOLHUKAM", "KEMENKO BIDANG HUKUM, HAM, IMIGRASI DAN PEMASYARAKATAN",
        "KEMENDAGRI", "KEMENLU", "SETKAB", "KKP", "KEMENKUM", "KEMENKEU",
        "SETNEG", "BAPPENAS", "KEMENHAN"
    ]

    new_doc = Document()
    new_table = new_doc.add_table(rows=0, cols=9)
    new_table.style = 'Table Grid'

    dim_counter = 14 # Starting after existing blocks
    
    # Process all rows
    for i, src_row in enumerate(table_src.rows):
        src_cells = src_row.cells
        
        if i < 155: # Header + Existing Blocks
            new_row = new_table.add_row()
            for j in range(9):
                copy_cell_xml(src_cells[j], new_row.cells[j])
        else: # Duplication needed
            content_check = src_cells[1].text.strip() or src_cells[4].text.strip()
            if content_check:
                dim_counter += 1
                for idx, agency in enumerate(agencies):
                    new_row = new_table.add_row()
                    # Col 0: No Dim
                    if idx == 0:
                        new_row.cells[0].text = str(dim_counter)
                    # Col 1-5: Content
                    for j in range(1, 6):
                        copy_cell_xml(src_cells[j], new_row.cells[j])
                    # Col 8: Agency
                    new_row.cells[8].text = agency
            else:
                new_row = new_table.add_row()
                for j in range(9):
                    copy_cell_xml(src_cells[j], new_row.cells[j])
        
        if i % 100 == 0:
            print(f"  Processed {i} source rows...")

    print("Phase 2: Applying Vertical Merges in blocks of 11...")
    # Apply merge to columns 0 to 5 (Index 0-5)
    # Starting from Row 2 (Index 1)
    for i in range(1, len(new_table.rows)):
        # Calculate position in 11-row block
        # Blocks start at 1, 12, 23, etc.
        pos = (i - 1) % 11
        merge_type = 'restart' if pos == 0 else 'continue'
        
        for col_idx in range(6):
            set_cell_vmerge(new_table.rows[i].cells[col_idx], merge_type)
            
    print(f"Saving final document: {output_path}")
    new_doc.save(output_path)
    print("Success! Formatting preserved and Vertical Merge fixed.")

if __name__ == "__main__":
    src = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026.docx"
    out = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_FINAL_FIXED.docx"
    rebuild_and_merge_fixed(src, out)
