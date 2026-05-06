import sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Add project root to sys.path
sys.path.insert(0, '/home/aseps/MCP/mcp-unified')

def set_vertical_merge(cell, merge_type):
    """
    merge_type: 'restart' or 'continue'
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vMerge = OxmlElement('w:vMerge')
    if merge_type == 'restart':
        vMerge.set(qn('w:val'), 'restart')
    tcPr.append(vMerge)

def merge_duplicate_rows_fast(file_path, output_path):
    print(f"Loading document for merging: {file_path}")
    doc = Document(file_path)
    if not doc.tables:
        return

    table = doc.tables[0]
    num_rows = len(table.rows)
    num_cols = 6 # We only merge first 6 columns (index 0-5)
    
    print(f"Analyzing {num_rows} rows for vertical merges...")
    
    for col_idx in range(num_cols):
        print(f"  Processing Column {col_idx + 1}...")
        current_val = None
        start_idx = -1
        
        for row_idx in range(num_rows):
            cell = table.cell(row_idx, col_idx)
            text = cell.text.strip()
            
            if text and text == current_val:
                # Same as previous, mark as continue
                set_vertical_merge(cell, 'continue')
            else:
                # New value or empty, restart merge here
                current_val = text
                if text: # Only merge if not empty
                    set_vertical_merge(cell, 'restart')
            
            if row_idx % 1000 == 0:
                print(f"    Processed {row_idx} rows in current column...")

    print(f"Saving merged document to: {output_path}")
    doc.save(output_path)
    print("Success! Table cells merged vertically.")

if __name__ == "__main__":
    input_file = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_COMPLETED.docx"
    output_file = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MERGED.docx"
    merge_duplicate_rows_fast(input_file, output_file)
