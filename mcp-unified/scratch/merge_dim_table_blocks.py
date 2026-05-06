import sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Add project root to sys.path
sys.path.insert(0, '/home/aseps/MCP/mcp-unified')

def set_vmerge_xml(tc, merge_type):
    tcPr = tc.get_or_add_tcPr()
    vMerge_list = tcPr.xpath('w:vMerge')
    if vMerge_list:
        vMerge = vMerge_list[0]
    else:
        vMerge = OxmlElement('w:vMerge')
        tcPr.append(vMerge)
    
    if merge_type == 'restart':
        vMerge.set(qn('w:val'), 'restart')
    else:
        if qn('w:val') in vMerge.attrib:
            del vMerge.attrib[qn('w:val')]

def merge_by_block_size(file_path, output_path, block_size=11):
    print(f"Loading document for block-based merging: {file_path}")
    doc = Document(file_path)
    if not doc.tables:
        return

    table = doc.tables[0]
    num_cols = 6
    rows = table._tbl.xpath('.//w:tr')
    total_rows = len(rows)
    
    print(f"Applying vertical merges in blocks of {block_size} for {total_rows} rows...")
    
    # Data starts from index 1 (Row 2), skip header at index 0
    for i in range(1, total_rows):
        # Calculate position within the current block
        # (i-1) % block_size == 0 means it's the first row of a block
        position_in_block = (i - 1) % block_size
        
        tr = rows[i]
        tcs = tr.xpath('.//w:tc')
        
        for j in range(min(num_cols, len(tcs))):
            tc = tcs[j]
            if position_in_block == 0:
                set_vmerge_xml(tc, 'restart')
            else:
                set_vmerge_xml(tc, 'continue')
        
        if i > 0 and i % 500 == 0:
            print(f"  Processed {i} rows...")

    print(f"Saving merged document to: {output_path}")
    doc.save(output_path)
    print("Success! Table merged strictly by blocks of 11.")

if __name__ == "__main__":
    # Use the _COMPLETED version as source to avoid double-merging artifacts
    input_file = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_COMPLETED.docx"
    output_file = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MERGED_FINAL.docx"
    merge_by_block_size(input_file, output_file, block_size=11)
