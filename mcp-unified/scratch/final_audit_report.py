from docx import Document

def final_audit_report():
    setkab_path = "/home/aseps/MCP/storage/office/DIM RUU Daerah Kepulauan_Setkab.docx"
    master_path = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MASTER_OFFICIAL_FINAL1.docx"
    
    def get_text(tc):
        return "".join([node.text for node in tc.xpath('.//w:t') if node.text]).strip()

    s_doc = Document(setkab_path)
    m_doc = Document(master_path)
    
    s_table = s_doc.tables[0]
    m_table = m_doc.tables[0]
    
    print("| No | Setkab DIM | Master DIM | Setkab Sub (C2) | Master Sub (C1) | Master Change (C2) |")
    print("|----|------------|------------|-----------------|-----------------|--------------------|")
    
    # Let's manually pick some rows that represent items
    # For Master, we look for rows with content in C0
    m_items = []
    for i, tr in enumerate(m_table._tbl.xpath('.//w:tr')):
        cells = tr.xpath('.//w:tc')
        if len(cells) >= 3:
            dim = get_text(cells[0])
            if dim and dim != "NO. DIM":
                m_items.append({
                    'dim': dim,
                    'c1': get_text(cells[1]),
                    'c2': get_text(cells[2])
                })
        if len(m_items) >= 10: break

    # For Setkab
    s_items = []
    for i, tr in enumerate(s_table._tbl.xpath('.//w:tr')):
        cells = tr.xpath('.//w:tc')
        if len(cells) >= 3:
            dim = get_text(cells[0])
            if dim and dim != "NO. DIM":
                s_items.append({
                    'dim': dim,
                    'c2': get_text(cells[2])
                })
        if len(s_items) >= 10: break

    limit = min(len(s_items), len(m_items))
    for i in range(limit):
        s = s_items[i]
        m = m_items[i]
        print(f"| {i+1} | {s['dim']} | {m['dim']} | {s['c2'][:20]}... | {m['c1'][:20]}... | {m['c2'][:20]}... |")

if __name__ == "__main__":
    final_audit_report()
