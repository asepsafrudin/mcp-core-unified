import os
import sys
import copy
from docx import Document

# Add project root to sys.path
sys.path.insert(0, '/home/aseps/MCP/mcp-unified')

def rebuild_dim_built_in_merge(input_path, output_path):
    print(f"Opening template: {input_path}")
    doc = Document(input_path)
    table = doc.tables[0]
    tbl_element = table._tbl
    
    # Cache source rows
    original_rows = list(table.rows)
    total_original = len(original_rows)
    
    agencies = [
        "KEMENKOPOLHUKAM", "KEMENKO BIDANG HUKUM, HAM, IMIGRASI DAN PEMASYARAKATAN",
        "KEMENDAGRI", "KEMENLU", "SETKAB", "KKP", "KEMENKUM", "KEMENKEU",
        "SETNEG", "BAPPENAS", "KEMENHAN"
    ]

    print("Step 1: Duplicating rows with formatting...")
    # Process duplication rows (from 155 onwards)
    dim_counter = 14
    for i in range(155, total_original):
        src_row = original_rows[i]
        content = src_row.cells[1].text.strip() or src_row.cells[4].text.strip()
        
        if content:
            dim_counter += 1
            for idx, agency in enumerate(agencies):
                new_row = table.add_row()
                
                # Clean Merge Strategy: Only fill content for the first row of the block (idx == 0)
                if idx == 0:
                    # Set NO. DIM
                    new_row.cells[0].text = str(dim_counter)
                    
                    # Copy formatting Col 1-5 (Deep Copy via XML for paragraphs)
                    for j in range(1, 6):
                        # Clear default empty paragraph in the new cell
                        for p in new_row.cells[j].paragraphs:
                            p._element.getparent().remove(p._element)
                        # Copy all paragraphs from source row to the new cell
                        for p in src_row.cells[j].paragraphs:
                            new_row.cells[j]._tc.append(copy.deepcopy(p._p))
                else:
                    # For subsequent rows in the block, keep columns 0-5 empty
                    new_row.cells[0].text = ""
                    # No content copied for cells[1] through cells[5]
                
                # Always set Agency name for Column 8 (as this column is NOT merged)
                new_row.cells[8].text = agency
        
        if i % 50 == 0:
            print(f"  Processed {i} source rows...")

    # Now we need to delete the original source rows (155 to end)
    # Important: delete from bottom up
    print("Deleting original source rows...")
    for i in range(total_original - 1, 154, -1):
        tr = original_rows[i]._tr
        tbl_element.remove(tr)

    print("Step 2: Applying Built-in Vertical Merge (Optimized)...")
    # All rows are now in the table. 
    # Existing blocks (1-154) + New blocks.
    all_rows = table.rows
    total_final_rows = len(all_rows)
    
    # Merger per block of 11
    # Skip header
    for i in range(1, total_final_rows, 11):
        # Merge columns 0 to 5 for rows i to i+10
        end_row_idx = min(i + 10, total_final_rows - 1)
        if end_row_idx > i:
            for col_idx in range(6):
                # Standard python-docx merge
                start_cell = all_rows[i].cells[col_idx]
                end_cell = all_rows[end_row_idx].cells[col_idx]
                start_cell.merge(end_cell)
        
        if i % 500 == 0:
            print(f"  Merged {i} rows...")

    print(f"Saving final verified document: {output_path}")
    doc.save(output_path)
    print("Success! Perfect DIM document with Built-in Merger.")

if __name__ == "__main__":
    src = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026.docx"
    out = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MASTER_OFFICIAL_FINAL.docx"
    rebuild_dim_built_in_merge(src, out)
