import os
import sys
from docx import Document
from docx.shared import Inches

# Add project root to sys.path
sys.path.insert(0, '/home/aseps/MCP/mcp-unified')

def complete_dim_table(input_path, output_path):
    print(f"Loading document: {input_path}")
    doc = Document(input_path)
    
    if not doc.tables:
        print("No tables found.")
        return

    original_table = doc.tables[0]
    total_rows = len(original_table.rows)
    print(f"Original rows: {total_rows}")

    agencies = [
        "KEMENKOPOLHUKAM",
        "KEMENKO BIDANG HUKUM, HAM, IMIGRASI DAN PEMASYARAKATAN",
        "KEMENDAGRI",
        "KEMENLU",
        "SETKAB",
        "KKP",
        "KEMENKUM",
        "KEMENKEU",
        "SETNEG",
        "BAPPENAS",
        "KEMENHAN"
    ]

    new_doc = Document()
    # Create new table with same number of columns
    new_table = new_doc.add_table(rows=0, cols=9)
    new_table.style = original_table.style

    # Copy header and first 155 rows (index 0 to 154)
    print("Copying first 155 rows...")
    for i in range(155):
        if i >= total_rows:
            break
        original_row = original_table.rows[i]
        new_row = new_table.add_row()
        for j in range(9):
            new_row.cells[j].text = original_row.cells[j].text

    # Process rows 156 to end (index 155 to total_rows-1)
    print("Processing and duplicating rows from 156 onwards...")
    for i in range(155, total_rows):
        original_row = original_table.rows[i]
        
        # Check if row has any content in Col 2 or Col 5
        content_check = original_row.cells[1].text.strip() or original_row.cells[4].text.strip()
        
        if content_check:
            # Duplicate this row 11 times
            for agency in agencies:
                new_row = new_table.add_row()
                # Copy first 6 columns
                for j in range(6):
                    new_row.cells[j].text = original_row.cells[j].text
                # Set agency in Col 9
                new_row.cells[8].text = agency
                # Col 7 and 8 remain empty or copied if original had something (usually empty)
                # For safety, let's copy original Col 7-8 if they have data
                new_row.cells[6].text = original_row.cells[6].text
                new_row.cells[7].text = original_row.cells[7].text
        else:
            # If empty row, just copy once
            new_row = new_table.add_row()
            for j in range(9):
                new_row.cells[j].text = original_row.cells[j].text

    print(f"Saving completed document: {output_path}")
    new_doc.save(output_path)
    print("Process complete!")

if __name__ == "__main__":
    input_file = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026.docx"
    output_file = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_COMPLETED.docx"
    complete_dim_table(input_file, output_file)
