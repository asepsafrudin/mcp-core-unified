import os
import sys
from docx import Document

# Add project root to sys.path
sys.path.insert(0, '/home/aseps/MCP/mcp-unified')

def complete_dim_table_super_fast(input_path, output_path):
    print(f"Reading original document: {input_path}")
    doc = Document(input_path)
    if not doc.tables:
        return
    
    original_table = doc.tables[0]
    original_rows = []
    for row in original_table.rows:
        original_rows.append([cell.text for cell in row.cells])
    
    total_original = len(original_rows)
    print(f"Read {total_original} rows.")

    agencies = [
        "KEMENKOPOLHUKAM",
        "KEMENKO BIDANG HUKUM, HAM, IMIGRASI DAN PEMASYARAKATAN",
        "KEMENDAGRI",
        "KEMENLU",
        "SETKAB",
        "KKP",
        "KEMENKUM",
        "KEMENKEU",
        "SETNEG",
        "BAPPENAS",
        "KEMENHAN"
    ]

    final_data = []
    # Copy first 155 rows
    for i in range(min(155, total_original)):
        final_data.append(original_rows[i])
    
    # Process from index 155
    print("Processing duplications in memory...")
    for i in range(155, total_original):
        row_data = original_rows[i]
        content_check = row_data[1].strip() or row_data[4].strip()
        
        if content_check:
            for agency in agencies:
                new_row = list(row_data) # Copy
                new_row[8] = agency # Set agency
                final_data.append(new_row)
        else:
            final_data.append(row_data)

    print(f"Total rows to be created: {len(final_data)}")

    # Create new document
    new_doc = Document()
    print("Pre-allocating table...")
    new_table = new_doc.add_table(rows=len(final_data), cols=9)
    new_table.style = 'Table Grid'

    print("Populating rows efficiently...")
    for i, row in enumerate(new_table.rows):
        row_data = final_data[i]
        for j, cell_text in enumerate(row_data):
            row.cells[j].text = cell_text
        if i > 0 and i % 500 == 0:
            print(f"  Populated {i} rows...")

    print(f"Saving to: {output_path}")
    new_doc.save(output_path)
    print("Success!")

if __name__ == "__main__":
    input_file = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026.docx"
    output_file = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_COMPLETED.docx"
    complete_dim_table_super_fast(input_file, output_file)
