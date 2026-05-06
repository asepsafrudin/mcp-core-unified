import os
import sys
import copy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Add project root to sys.path
sys.path.insert(0, '/home/aseps/MCP/mcp-unified')

def copy_paragraphs_xml(src_tc, dst_tc):
    """Deep copy paragraphs from src tc to dst tc."""
    # Clear dst_tc
    for p in dst_tc.xpath('w:p'):
        dst_tc.remove(p)
    # Copy from src
    for p in src_tc.xpath('w:p'):
        dst_tc.append(copy.deepcopy(p))

def set_vmerge_xml(tc, merge_type):
    tcPr = tc.get_or_add_tcPr()
    vMerge = tcPr.xpath('w:vMerge')
    if not vMerge:
        vMerge = OxmlElement('w:vMerge')
        tcPr.append(vMerge)
    else:
        vMerge = vMerge[0]
    
    if merge_type == 'restart':
        vMerge.set(qn('w:val'), 'restart')
    else:
        if qn('w:val') in vMerge.attrib:
            del vMerge.attrib[qn('w:val')]

def rebuild_and_merge_ultra_fast(input_path, output_path):
    print(f"Loading master document: {input_path}")
    doc_src = Document(input_path)
    table_src = doc_src.tables[0]
    
    # Store source cells data in memory
    print("Reading source data into memory...")
    source_data = []
    for row in table_src.rows:
        source_data.append([cell._tc for cell in row.cells])
    
    agencies = [
        "KEMENKOPOLHUKAM", "KEMENKO BIDANG HUKUM, HAM, IMIGRASI DAN PEMASYARAKATAN",
        "KEMENDAGRI", "KEMENLU", "SETKAB", "KKP", "KEMENKUM", "KEMENKEU",
        "SETNEG", "BAPPENAS", "KEMENHAN"
    ]

    # Prepare final data list
    final_cells_xml = []
    
    print("Preparing duplicated row data...")
    # Header
    final_cells_xml.append(source_data[0])
    
    # Existing blocks (1-154)
    for i in range(1, 155):
        final_cells_xml.append(source_data[i])
        
    # Duplication part (155 to end)
    dim_counter = 14
    for i in range(155, len(source_data)):
        src_row_tc = source_data[i]
        # Check content in Col 2 (index 1) or Col 5 (index 4)
        content_p = src_row_tc[1].xpath('.//w:t') or src_row_tc[4].xpath('.//w:t')
        
        if content_p:
            dim_counter += 1
            for idx, agency in enumerate(agencies):
                new_row_tc_data = [copy.deepcopy(tc) for tc in src_row_tc]
                
                # Set NO. DIM
                if idx == 0:
                    # Clear and set DIM number
                    no_dim_p = new_row_tc_data[0].xpath('.//w:p')[0]
                    for t in no_dim_p.xpath('.//w:t'):
                        t.text = str(dim_counter)
                else:
                    # Clear text for continue cells in Col 0
                    for t in new_row_tc_data[0].xpath('.//w:t'):
                        t.text = ""

                # Set Agency in Col 9 (index 8)
                agency_p = new_row_tc_data[8].xpath('.//w:p')[0]
                for t in agency_p.xpath('.//w:t'):
                    t.text = agency
                
                # Clear Col 7, 8 (index 6, 7)
                for t in new_row_tc_data[6].xpath('.//w:t'): t.text = ""
                for t in new_row_tc_data[7].xpath('.//w:t'): t.text = ""
                
                # Set vMerge markers
                merge_type = 'restart' if idx == 0 else 'continue'
                for col_idx in range(6):
                    set_vmerge_xml(new_row_tc_data[col_idx], merge_type)
                
                final_cells_xml.append(new_row_tc_data)
        else:
            final_cells_xml.append(source_data[i])

    print(f"Total rows to create: {len(final_cells_xml)}")
    
    new_doc = Document()
    new_table = new_doc.add_table(rows=len(final_cells_xml), cols=9)
    new_table.style = 'Table Grid'
    
    print("Injecting XML rows into new table...")
    # Replace table body with our custom built rows
    tbl = new_table._tbl
    # Remove existing rows
    for tr in tbl.xpath('w:tr'):
        tbl.remove(tr)
    
    for i, row_tc_list in enumerate(final_cells_xml):
        tr = OxmlElement('w:tr')
        for tc in row_tc_list:
            tr.append(tc)
        tbl.append(tr)
        if i % 1000 == 0:
            print(f"  Injected {i} rows...")

    print(f"Saving final document: {output_path}")
    new_doc.save(output_path)
    print("Success! Formatting preserved and Vertical Merge completed ultra-fast.")

if __name__ == "__main__":
    src = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026.docx"
    out = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_FINAL_PERFECT.docx"
    rebuild_and_merge_ultra_fast(src, out)
