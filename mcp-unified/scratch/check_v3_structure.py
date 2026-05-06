from docx import Document
import json
import os

def check_table_structure():
    doc_path = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MASTER_OFFICIAL_FINAL_KEMENKEU_V3.docx"
    doc = Document(doc_path)
    if not doc.tables:
        print("No tables found")
        return
    
    table = doc.tables[0]
    print(f"Table has {len(table.rows)} rows and {len(table.columns)} columns.")
    
    # Read headers (row 0)
    headers = [cell.text.strip() for cell in table.rows[0].cells]
    print("Headers:", headers)
    
    # Sample row 1
    if len(table.rows) > 1:
        sample = [cell.text.strip() for cell in table.rows[1].cells]
        print("Sample Row 1:", sample)

if __name__ == "__main__":
    check_table_structure()
