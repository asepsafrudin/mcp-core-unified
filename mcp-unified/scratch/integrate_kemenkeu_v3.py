import pdfplumber
from docx import Document
from docx.shared import Pt
import re

def normalize(text):
    if not text: return ""
    return re.sub(r'[^a-zA-Z0-9]', '', text).lower()

def extract_pdf_clean(pdf_path):
    all_data = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if not text: continue
            
            # Find rows using pattern "Number Pasal X"
            # Pattern: \n(\d+)\s+(Pasal\s+\d+)
            matches = list(re.finditer(r'\n(\d+)\s+(Pasal\s+\d+)', text))
            for i, match in enumerate(matches):
                no = match.group(1)
                pasal = match.group(2)
                start_pos = match.start()
                end_pos = matches[i+1].start() if i+1 < len(matches) else len(text)
                
                block = text[start_pos:end_pos]
                # Try to separate columns roughly based on common PDF patterns
                # This is still fuzzy but better than fixed coordinates
                lines = block.split('\n')
                muatan = " ".join(lines[1:5]) # Heuristic for muatan text
                
                all_data.append({
                    'no': no,
                    'pasal': pasal,
                    'muatan': muatan.strip(),
                    'full_block': block
                })
    return all_data

def integrate_v3():
    pdf_path = "/home/aseps/MCP/storage/office/raw/S-12.PK.PK1.2026 Reviu DIM RUU Kepulauan - DJPK-tabel only.pdf"
    master_path = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MASTER_OFFICIAL_FINAL_MIGRATED.docx"
    output_path = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MASTER_OFFICIAL_FINAL_KEMENKEU_V3.docx"
    
    # 1. Parse PDF
    pdf_items = extract_pdf_clean(pdf_path)
    print(f"Extracted {len(pdf_items)} items from PDF.")
    
    # 2. Open Master
    print(f"Opening Master: {master_path}")
    doc = Document(master_path)
    table = doc.tables[0]
    master_rows = list(table.rows)
    
    # Index Master (starting from Row 1)
    master_index = []
    for i, row in enumerate(master_rows):
        if i == 0: continue
        txt = row.cells[1].text.strip()
        master_index.append((i, txt, normalize(txt)))
    
    updated_count = 0
    failed_items = []
    
    print("Processing updates from Row 1...")
    for p_item in pdf_items:
        p_norm = normalize(p_item['muatan'])
        pasal_num = p_item['pasal'] # e.g. "Pasal 4"
        
        match_idx = -1
        # Search for best match
        for m_idx, m_txt, m_norm in master_index:
            # Check if Pasal match + partial substance match
            if pasal_num in m_txt:
                # If it's a Pasal row, or the next few rows match the substance
                match_found = False
                if len(m_norm) > 5 and (m_norm in p_norm or p_norm in m_norm):
                    match_found = True
                else:
                    # Look at next few rows for substance match
                    for offset in range(1, 10):
                        next_idx = m_idx + offset
                        if next_idx < len(master_rows):
                            next_norm = normalize(master_rows[next_idx].cells[1].text)
                            if next_norm and (next_norm in p_norm or p_norm in next_norm):
                                match_found = True
                                break
                
                if match_found:
                    # Find KEMENKEU row in this block
                    for k_offset in range(-5, 15):
                        target_idx = m_idx + k_offset
                        if 0 < target_idx < len(master_rows):
                            agency = master_rows[target_idx].cells[5].text.strip()
                            if agency == "KEMENKEU":
                                # Update (Simulating extraction for now as full parsing is complex)
                                # We'll put the block as Keterangan for now to see success
                                master_rows[target_idx].cells[7].text = f"KEMENKEU REVIU: {p_item['full_block'][:500]}"
                                
                                # Format Arial 10
                                for paragraph in master_rows[target_idx].cells[7].paragraphs:
                                    for run in paragraph.runs:
                                        run.font.name = 'Arial'
                                        run.font.size = Pt(10)
                                        
                                updated_count += 1
                                match_idx = m_idx
                                break
                if match_idx != -1: break
        
        if match_idx == -1:
            failed_items.append(p_item)

    print(f"\nSummary:")
    print(f"Success: {updated_count}")
    print(f"Failed: {len(failed_items)}")
    
    if failed_items:
        print("\nFailed Items List:")
        for f in failed_items:
            print(f"  PDF No. {f['no']} | {f['pasal']}")
            
    doc.save(output_path)
    print(f"Result saved to: {output_path}")

if __name__ == "__main__":
    integrate_v3()
