import sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Add project root to sys.path
sys.path.insert(0, '/home/aseps/MCP/mcp-unified')

def set_vertical_merge(cell, merge_type):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vMerge = OxmlElement('w:vMerge')
    if merge_type == 'restart':
        vMerge.set(qn('w:val'), 'restart')
    tcPr.append(vMerge)

def merge_duplicate_rows_optimized(file_path, output_path):
    print(f"Loading document for merging: {file_path}")
    doc = Document(file_path)
    if not doc.tables:
        return

    table = doc.tables[0]
    num_cols = 6
    
    # Store previous values to detect duplicates
    prev_values = [None] * num_cols
    
    print(f"Applying vertical merges to {len(table.rows)} rows...")
    for i, row in enumerate(table.rows):
        for j in range(num_cols):
            cell = row.cells[j]
            text = cell.text.strip()
            
            if text and text == prev_values[j]:
                # Mark as continue
                set_vertical_merge(cell, 'continue')
            else:
                # Mark as restart
                prev_values[j] = text
                if text:
                    set_vertical_merge(cell, 'restart')
        
        if i > 0 and i % 500 == 0:
            print(f"  Processed {i} rows...")

    print(f"Saving merged document to: {output_path}")
    doc.save(output_path)
    print("Success!")

if __name__ == "__main__":
    input_file = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_COMPLETED.docx"
    output_file = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MERGED.docx"
    merge_duplicate_rows_optimized(input_file, output_file)
