from docx import Document
import re

def normalize_dim(dim_text):
    return re.sub(r'[^0-9a-zA-Z]', '', dim_text).strip()

def get_setkab_dict(file_path):
    print(f"Reading Setkab data from {file_path}...")
    doc = Document(file_path)
    table = doc.tables[0]
    data_dict = {}
    
    for i, row in enumerate(table.rows):
        if i == 0: continue
        try:
            # Berdasarkan audit: DIM ada di Index 0, Usulan 2025 ada di Index 5
            dim_raw = row.cells[0].text.strip()
            # Usulan Perubahan dan DIM Tahun 2025 di Index 5
            sub_raw = row.cells[5].text.strip()
            
            dim_key = normalize_dim(dim_raw)
            if dim_key:
                data_dict[dim_key] = sub_raw
        except Exception:
            continue
    return data_dict

def update_master():
    setkab_path = "/home/aseps/MCP/storage/office/DIM RUU Daerah Kepulauan_Setkab.docx"
    master_path = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MASTER_OFFICIAL_FINAL1.docx"
    output_path = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MASTER_OFFICIAL_FINAL_UPDATED.docx"
    
    setkab_data = get_setkab_dict(setkab_path)
    print(f"Extracted {len(setkab_data)} items from Setkab.")
    
    print(f"Opening Master file: {master_path}")
    doc = Document(master_path)
    table = doc.tables[0]
    
    current_dim_key = ""
    updated_count = 0
    
    print("Processing Master rows...")
    for i, row in enumerate(table.rows):
        if i == 0: continue # Skip header
        
        # Update current DIM Key if present
        dim_text = row.cells[0].text.strip()
        if dim_text:
            current_dim_key = normalize_dim(dim_text)
        
        # Check if this is a SETKAB row (Col 8)
        # Based on previous audit, Col 8 contains the Agency name
        try:
            keterangan_cell = row.cells[8]
            if keterangan_cell.text.strip() == "SETKAB":
                # Look up in Setkab data
                if current_dim_key in setkab_data:
                    new_content = setkab_data[current_dim_key]
                    keterangan_cell.text = f"SETKAB: {new_content}"
                    updated_count += 1
        except Exception:
            continue

    print(f"Successfully updated {updated_count} SETKAB rows.")
    print(f"Saving updated document to: {output_path}")
    doc.save(output_path)
    print("Done!")

if __name__ == "__main__":
    update_master()
