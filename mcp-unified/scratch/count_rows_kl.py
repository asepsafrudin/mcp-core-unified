from docx import Document

def count_rows(file_path):
    doc = Document(file_path)
    if not doc.tables:
        print("No tables found.")
        return
    table = doc.tables[0]
    print(f"Total rows in {file_path}: {len(table.rows)}")

if __name__ == "__main__":
    count_rows("/home/aseps/MCP/storage/office/21042026 DIM RUU DAERAH KEPULAUAN GABUNG KL.docx")
