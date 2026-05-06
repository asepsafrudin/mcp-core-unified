import json
import os
from docx import Document

def convert_docx_to_json():
    doc_path = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MASTER_OFFICIAL_FINAL_KEMENKEU_V3.docx"
    output_dir=os.getenv("OUTPUT_DIR", "/home/aseps/MCP/storage/office/data" if not os.getenv("CI") else "DUMMY")
    os.makedirs(output_dir, exist_ok=True)
    
    print(f"Opening document: {doc_path}")
    doc = Document(doc_path)
    if not doc.tables:
        print("Error: No tables found in document.")
        return
    
    table = doc.tables[0]
    rows = list(table.rows)
    
    # Headers
    headers = [cell.text.strip() for cell in rows[0].cells]
    # Clean headers for JSON keys (remove newlines, replace space with underscore, lowercase)
    clean_headers = [h.replace("\n", " ").replace("  ", " ").strip().lower().replace(" ", "_") for h in headers]
    
    data = []
    print(f"Processing {len(rows)-1} rows...")
    
    for i, row in enumerate(rows[1:]):
        item = {}
        for j, cell in enumerate(row.cells):
            key = clean_headers[j] if j < len(clean_headers) else f"col_{j}"
            item[key] = cell.text.strip()
        data.append(item)
        
        if (i+1) % 500 == 0:
            print(f"  Processed {i+1} rows...")

    # Determine if chunking is needed
    # Let's check estimated size
    json_str = json.dumps(data, indent=2)
    size_mb = len(json_str) / (1024 * 1024)
    print(f"Estimated JSON size: {size_mb:.2f} MB")
    
    chunk_size = 1000
    if size_mb > 10: # If > 10MB, chunk it
        print("Size > 10MB, performing chunking...")
        for i in range(0, len(data), chunk_size):
            chunk = data[i : i + chunk_size]
            chunk_file = os.path.join(output_dir, f"dim_v3_chunk_{i//chunk_size + 1}.json")
            with open(chunk_file, 'w', encoding='utf-8') as f:
                json.dump(chunk, f, indent=2, ensure_ascii=False)
            print(f"  Saved chunk {i//chunk_size + 1} to {chunk_file}")
    else:
        output_file = os.path.join(output_dir, "dim_v3_full.json")
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(json_str)
        print(f"Saved full JSON to {output_file}")

if __name__ == "__main__":
    convert_docx_to_json()