import sys
from docx import Document
from docx.oxml.ns import qn

# Add project root to sys.path
sys.path.insert(0, '/home/aseps/MCP/mcp-unified')

def audit_perfect_document(file_path):
    print(f"Auditing perfect document: {file_path}\n")
    doc = Document(file_path)
    table = doc.tables[0]
    rows = table.rows
    total_rows = len(rows)
    
    print(f"--- General Statistics ---")
    print(f"Total Rows: {total_rows}")
    
    agencies = [
        "KEMENKOPOLHUKAM", "KEMENKO BIDANG HUKUM, HAM, IMIGRASI DAN PEMASYARAKATAN",
        "KEMENDAGRI", "KEMENLU", "SETKAB", "KKP", "KEMENKUM", "KEMENKEU",
        "SETNEG", "BAPPENAS", "KEMENHAN"
    ]
    
    print(f"\n--- Verification of Blocks ---")
    sample_blocks = [1, 10, 100, 300]
    for b_idx in sample_blocks:
        start_row = 1 + (b_idx - 1) * 11
        if start_row < total_rows:
            print(f"Block {b_idx} (Row {start_row + 1}):")
            no_dim = rows[start_row].cells[0].text.strip()
            agency_1 = rows[start_row].cells[8].text.strip()
            
            print(f"  NO. DIM: {no_dim}")
            print(f"  Agency: {agency_1}")
            
            # Check vMerge via XML
            tc = rows[start_row].cells[1]._tc
            vMerge = tc.get_or_add_tcPr().xpath('w:vMerge')
            merge_val = vMerge[0].get(qn('w:val')) if vMerge else "NONE"
            print(f"  Merge (Row {start_row + 1}): {merge_val}")
            
            tc_next = rows[min(start_row + 1, total_rows-1)].cells[1]._tc
            vMerge_next = tc_next.get_or_add_tcPr().xpath('w:vMerge')
            merge_val_next = "continue" if vMerge_next and vMerge_next[0].get(qn('w:val')) is None else "FAIL"
            print(f"  Merge (Row {start_row + 2}): {merge_val_next}")

    print(f"\n--- Formatting Preservation Check ---")
    p_count = len(rows[1].cells[1].paragraphs)
    print(f"Paragraphs in Row 2 Col 2: {p_count}")

if __name__ == "__main__":
    audit_perfect_document("/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_PERFECT_AUDITED.docx")
