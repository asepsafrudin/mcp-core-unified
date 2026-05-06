import sys
from docx import Document
from docx.oxml.ns import qn

# Add project root to sys.path
sys.path.insert(0, '/home/aseps/MCP/mcp-unified')

def audit_final_document(file_path):
    print(f"Auditing final document: {file_path}\n")
    doc = Document(file_path)
    table = doc.tables[0]
    rows = table.rows
    total_rows = len(rows)
    
    print(f"--- General Statistics ---")
    print(f"Total Rows: {total_rows}")
    
    agencies = [
        "KEMENKOPOLHUKAM", "KEMENKO BIDANG HUKUM, HAM, IMIGRASI DAN PEMASYARAKATAN",
        "KEMENDAGRI", "KEMENLU", "SETKAB", "KKP", "KEMENKUM", "KEMENKEU",
        "SETNEG", "BAPPENAS", "KEMENHAN"
    ]
    
    print(f"\n--- Verification of Blocks (Every 11 rows) ---")
    # Check a few blocks
    sample_blocks = [1, 10, 100, 300]
    for b_idx in sample_blocks:
        start_row = 1 + (b_idx - 1) * 11
        if start_row < total_rows:
            print(f"Block {b_idx} (Starting at row {start_row + 1}):")
            # Check No Dim
            no_dim = rows[start_row].cells[0].text.strip()
            # Check Agency in the 1st and 11th row of the block
            agency_1 = rows[start_row].cells[8].text.strip()
            agency_11 = rows[min(start_row + 10, total_rows-1)].cells[8].text.strip()
            
            print(f"  NO. DIM: {no_dim}")
            print(f"  Agency 1: {agency_1} (Expected: {agencies[0]})")
            print(f"  Agency 11: {agency_11} (Expected: {agencies[10]})")
            
            # Check Merging (vMerge)
            tc = rows[start_row].cells[1]._tc
            tcPr = tc.get_or_add_tcPr()
            vMerge = tcPr.xpath('w:vMerge')
            merge_status = vMerge[0].get(qn('w:val')) if vMerge else "NONE"
            print(f"  Col 2 Merge Status (Row {start_row + 1}): {merge_status}")
            
            tc_cont = rows[min(start_row + 1, total_rows-1)].cells[1]._tc
            tcPr_cont = tc_cont.get_or_add_tcPr()
            vMerge_cont = tcPr_cont.xpath('w:vMerge')
            merge_status_cont = "continue" if vMerge_cont and vMerge_cont[0].get(qn('w:val')) is None else "ERROR"
            print(f"  Col 2 Merge Status (Row {start_row + 2}): {merge_status_cont}")

    print(f"\n--- Formatting Preservation Check ---")
    # Check paragraph count in Row 2 Col 2 (index 1, 1)
    p_count = len(rows[1].cells[1].paragraphs)
    print(f"Paragraphs in Row 2 Col 2: {p_count} (Original was 10)")
    
    if p_count == 10:
        print("Success: Multi-paragraph formatting is preserved.")
    else:
        print(f"Warning: Paragraph count is {p_count}, formatting might be lost.")

if __name__ == "__main__":
    audit_final_document("/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_STABLE_FINAL.docx")
