import pdfplumber
from docx import Document
from docx.shared import Pt
import re

def normalize(text):
    if not text: return ""
    return re.sub(r'[^a-zA-Z0-9]', '', text).lower()

def extract_pdf_v4(pdf_path):
    print(f"Extracting PDF with high precision: {pdf_path}")
    all_data = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            table = page.extract_table()
            if table:
                for row in table:
                    if not row or not row[0] or not row[0].isdigit(): continue
                    all_data.append({
                        'no': row[0].strip(),
                        'pasal': row[1].strip() if row[1] else "",
                        'muatan': row[1].strip() if row[1] else "", # Muatan is often in Col 1 or 2
                        'tanggapan': row[2].strip() if row[2] else "",
                        'usulan': row[3].strip() if row[3] else "",
                        'keterangan': row[4].strip() if row[4] else ""
                    })
    return all_data

def integrate_v4():
    pdf_path = "/home/aseps/MCP/storage/office/raw/S-12.PK.PK1.2026 Reviu DIM RUU Kepulauan - DJPK-tabel only.pdf"
    master_path = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MASTER_OFFICIAL_FINAL_MIGRATED.docx"
    output_path = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MASTER_OFFICIAL_FINAL_KEMENKEU_V4.docx"
    
    pdf_items = extract_pdf_v4(pdf_path)
    print(f"Total PDF items: {len(pdf_items)}")
    
    doc = Document(master_path)
    table = doc.tables[0]
    master_rows = list(table.rows)
    
    updated_count = 0
    
    for p_item in pdf_items:
        # Search substance in PDF muatan text
        p_muatan = p_item['muatan']
        p_norm = normalize(p_muatan)
        
        match_found = False
        for i in range(1, len(master_rows)):
            m_txt = master_rows[i].cells[1].text.strip()
            m_norm = normalize(m_txt)
            
            # Smart matching: check if Master substance is part of PDF Muatan or vice versa
            if len(m_norm) > 15 and (m_norm in p_norm or p_norm in m_norm):
                # Found block! Find KEMENKEU row in this block
                for offset in range(-5, 11):
                    target_idx = i + offset
                    if 0 < target_idx < len(master_rows):
                        agency = master_rows[target_idx].cells[5].text.strip()
                        if agency == "KEMENKEU":
                            combined = f"KEMENKEU: {p_item['tanggapan']}"
                            if p_item['usulan']: combined += f" | Usulan: {p_item['usulan']}"
                            if p_item['keterangan']: combined += f" | Ket: {p_item['keterangan']}"
                            
                            master_rows[target_idx].cells[7].text = combined
                            # Format Arial 10
                            for p in master_rows[target_idx].cells[7].paragraphs:
                                for r in p.runs:
                                    r.font.name = 'Arial'
                                    r.font.size = Pt(10)
                            
                            updated_count += 1
                            match_found = True
                            break
                if match_found: break
        
        if not match_found:
            print(f"  FAILED: No match for Item {p_item['no']} ({p_item['muatan'][:40]}...)")

    print(f"\nFinal Summary V4:")
    print(f"Successfully integrated {updated_count} items.")
    doc.save(output_path)
    print(f"Saved to: {output_path}")

if __name__ == "__main__":
    integrate_v4()
