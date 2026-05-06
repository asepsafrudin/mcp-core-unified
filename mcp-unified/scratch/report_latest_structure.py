from docx import Document

def report_latest_structure(file_path):
    print(f"Reporting Latest Structure for: {file_path}")
    doc = Document(file_path)
    if not doc.tables:
        print("No tables found.")
        return
        
    table = doc.tables[0]
    
    # Inspect Header (Row 0)
    print("\n--- STRUKTUR KOLOM TERBARU ---")
    header_row = table.rows[0]
    for i, cell in enumerate(header_row.cells):
        print(f"Kolom {i}: {cell.text.strip().replace('\n', ' ')}")
    
    # Check Font on a sample cell (e.g., Row 1, Col 1)
    print("\n--- VERIFIKASI FONT (Sampel Baris 1, Kolom 1) ---")
    sample_cell = table.rows[1].cells[1]
    if sample_cell.paragraphs:
        p = sample_cell.paragraphs[0]
        if p.runs:
            run = p.runs[0]
            print(f"Jenis Font: {run.font.name}")
            print(f"Ukuran Font: {run.font.size.pt if run.font.size else 'Default'}")

if __name__ == "__main__":
    target = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MASTER_OFFICIAL_FINAL_GABUNG_KL.docx"
    report_latest_structure(target)
