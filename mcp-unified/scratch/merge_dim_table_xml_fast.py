import sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Add project root to sys.path
sys.path.insert(0, '/home/aseps/MCP/mcp-unified')

def set_vmerge_xml(tc, merge_type):
    tcPr = tc.get_or_add_tcPr()
    vMerge_list = tcPr.xpath('w:vMerge')
    if vMerge_list:
        vMerge = vMerge_list[0]
    else:
        vMerge = OxmlElement('w:vMerge')
        tcPr.append(vMerge)
    
    if merge_type == 'restart':
        vMerge.set(qn('w:val'), 'restart')
    else:
        if qn('w:val') in vMerge.attrib:
            del vMerge.attrib[qn('w:val')]

def get_tc_text(tc):
    return "".join(t.text for t in tc.xpath('.//w:t') if t.text is not None)

def merge_duplicate_rows_xml_fast(file_path, output_path):
    print(f"Loading document for ultra-fast merging: {file_path}")
    doc = Document(file_path)
    if not doc.tables:
        return

    table = doc.tables[0]
    num_cols = 6
    
    prev_values = [None] * num_cols
    
    print("Applying vertical merges using direct XML path (Ultra Fast)...")
    # Get all row elements
    rows = table._tbl.xpath('.//w:tr')
    
    for i, tr in enumerate(rows):
        tcs = tr.xpath('.//w:tc')
        for j in range(min(num_cols, len(tcs))):
            tc = tcs[j]
            text = get_tc_text(tc).strip()
            
            if text and text == prev_values[j]:
                set_vmerge_xml(tc, 'continue')
            else:
                prev_values[j] = text
                if text:
                    set_vmerge_xml(tc, 'restart')
        
        if i > 0 and i % 1000 == 0:
            print(f"  Processed {i} rows...")

    print(f"Saving merged document to: {output_path}")
    doc.save(output_path)
    print("Success!")

if __name__ == "__main__":
    input_file = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_COMPLETED.docx"
    output_file = "/home/aseps/MCP/storage/office/R1 DIM RUU Kepualauan 13022026_MERGED.docx"
    merge_duplicate_rows_xml_fast(input_file, output_file)
