"""
Unit Tests — Document Converter Tool (Fase 1)

Menguji konversi DOCX/XLSX ke Clean Markdown,
defragmentasi XML split runs, dan mode sampling XLSX.
"""
import os
import sys
import tempfile
import shutil
import pytest
from pathlib import Path

# Pastikan mcp-unified bisa diimport
sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from tools.office.document_converter_tool import (
    convert_docx_to_markdown,
    convert_xlsx_to_markdown,
    defragment_docx_xml,
    _html_to_markdown,
)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def temp_dir():
    """Temporary directory untuk test files."""
    d = tempfile.mkdtemp()
    yield d
    shutil.rmtree(d, ignore_errors=True)


@pytest.fixture
def sample_docx(temp_dir):
    """Buat file DOCX sederhana untuk testing."""
    from docx import Document

    path = os.path.join(temp_dir, 'test_doc.docx')
    doc = Document()
    doc.core_properties.title = 'Test Document'
    doc.core_properties.author = 'Unit Test'

    doc.add_heading('Heading Satu', level=1)
    doc.add_paragraph('Ini adalah paragraf pertama dengan teks biasa.')
    doc.add_heading('Heading Dua', level=2)
    doc.add_paragraph('Paragraf kedua.')

    # Tabel
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = 'Nama'
    table.rows[0].cells[1].text = 'Nilai'
    table.rows[1].cells[0].text = 'Alpha'
    table.rows[1].cells[1].text = '100'
    table.rows[2].cells[0].text = 'Beta'
    table.rows[2].cells[1].text = '90'

    doc.save(path)
    return path


@pytest.fixture
def sample_xlsx(temp_dir):
    """Buat file XLSX sederhana untuk testing."""
    from openpyxl import Workbook

    path = os.path.join(temp_dir, 'test_sheet.xlsx')
    wb = Workbook()
    ws = wb.active
    ws.title = 'DataSheet'

    # Header + 5 baris data
    ws.append(['ID', 'Nama', 'Skor'])
    for i in range(1, 6):
        ws.append([i, f'Item-{i}', 80 + i])

    wb.save(path)
    return path


@pytest.fixture
def large_xlsx(temp_dir):
    """Buat file XLSX besar (250 baris) untuk test sampling."""
    from openpyxl import Workbook

    path = os.path.join(temp_dir, 'large_sheet.xlsx')
    wb = Workbook()
    ws = wb.active
    ws.title = 'BigData'

    ws.append(['ID', 'Value', 'Category'])
    for i in range(1, 251):
        ws.append([i, i * 10, f'Cat-{i % 5}'])

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Tests: HTML → Markdown internal converter
# ---------------------------------------------------------------------------

class TestHtmlToMarkdown:
    def test_heading_conversion(self):
        html = '<h1>Title</h1><h2>Subtitle</h2>'
        md = _html_to_markdown(html)
        assert '# Title' in md
        assert '## Subtitle' in md

    def test_bold_italic(self):
        html = '<p><strong>bold</strong> and <em>italic</em></p>'
        md = _html_to_markdown(html)
        assert '**bold**' in md
        assert '*italic*' in md

    def test_table(self):
        html = '<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>'
        md = _html_to_markdown(html)
        assert '| A | B |' in md
        assert '| --- | --- |' in md
        assert '| 1 | 2 |' in md

    def test_unordered_list(self):
        html = '<ul><li>item one</li><li>item two</li></ul>'
        md = _html_to_markdown(html)
        assert '- item one' in md
        assert '- item two' in md


# ---------------------------------------------------------------------------
# Tests: DOCX → Markdown
# ---------------------------------------------------------------------------

class TestConvertDocxToMarkdown:
    def test_successful_conversion(self, sample_docx):
        result = convert_docx_to_markdown(sample_docx)
        assert result['success'] is True
        assert 'markdown' in result
        assert len(result['markdown']) > 0
        assert result['conversion_method'] == 'mammoth + html_to_markdown'

    def test_metadata_extraction(self, sample_docx):
        result = convert_docx_to_markdown(sample_docx)
        assert result['success'] is True
        assert result['metadata']['title'] == 'Test Document'
        assert result['metadata']['author'] == 'Unit Test'

    def test_statistics(self, sample_docx):
        result = convert_docx_to_markdown(sample_docx)
        assert result['success'] is True
        stats = result['statistics']
        assert stats['word_count'] > 0
        assert stats['character_count'] > 0

    def test_file_not_found(self):
        result = convert_docx_to_markdown('/nonexistent/file.docx')
        assert result['success'] is False
        assert 'not found' in result['error'].lower()

    def test_wrong_extension(self, temp_dir):
        path = os.path.join(temp_dir, 'test.txt')
        with open(path, 'w') as f:
            f.write('hello')
        result = convert_docx_to_markdown(path)
        assert result['success'] is False
        assert '.docx' in result['error']


# ---------------------------------------------------------------------------
# Tests: XLSX → Markdown
# ---------------------------------------------------------------------------

class TestConvertXlsxToMarkdown:
    def test_successful_conversion(self, sample_xlsx):
        result = convert_xlsx_to_markdown(sample_xlsx)
        assert result['success'] is True
        assert 'markdown' in result
        assert '| ID | Nama | Skor |' in result['markdown']

    def test_statistics(self, sample_xlsx):
        result = convert_xlsx_to_markdown(sample_xlsx)
        assert result['success'] is True
        stats = result['statistics']
        assert stats['total_rows'] == 5  # 5 data rows (header excluded)
        assert stats['total_columns'] == 3

    def test_no_truncation_for_small_data(self, sample_xlsx):
        result = convert_xlsx_to_markdown(sample_xlsx, max_rows=100)
        assert result['success'] is True
        assert result['sampling_info']['is_truncated'] is False
        assert result['requires_confirmation'] is False

    def test_head_sampling(self, large_xlsx):
        result = convert_xlsx_to_markdown(large_xlsx, max_rows=50, sampling_mode='head')
        assert result['success'] is True
        assert result['sampling_info']['is_truncated'] is True
        assert result['statistics']['displayed_rows'] == 50
        assert result['requires_confirmation'] is True

    def test_tail_sampling(self, large_xlsx):
        result = convert_xlsx_to_markdown(large_xlsx, max_rows=30, sampling_mode='tail')
        assert result['success'] is True
        assert result['statistics']['displayed_rows'] == 30

    def test_sample_mode(self, large_xlsx):
        result = convert_xlsx_to_markdown(large_xlsx, max_rows=20, sampling_mode='sample')
        assert result['success'] is True
        assert result['statistics']['displayed_rows'] == 20

    def test_chunk_mode(self, large_xlsx):
        result = convert_xlsx_to_markdown(
            large_xlsx, sampling_mode='chunk', chunk_index=0, chunk_size=50
        )
        assert result['success'] is True
        assert result['chunking_info'] is not None
        assert result['chunking_info']['current_chunk'] == 0
        assert result['chunking_info']['has_next'] is True

    def test_chunk_navigation(self, large_xlsx):
        # Chunk 0
        r0 = convert_xlsx_to_markdown(
            large_xlsx, sampling_mode='chunk', chunk_index=0, chunk_size=100
        )
        # Chunk 1
        r1 = convert_xlsx_to_markdown(
            large_xlsx, sampling_mode='chunk', chunk_index=1, chunk_size=100
        )
        # Chunk 2
        r2 = convert_xlsx_to_markdown(
            large_xlsx, sampling_mode='chunk', chunk_index=2, chunk_size=100
        )

        assert r0['chunking_info']['has_next'] is True
        assert r1['chunking_info']['has_next'] is True
        assert r2['chunking_info']['has_next'] is False

    def test_full_mode(self, large_xlsx):
        result = convert_xlsx_to_markdown(large_xlsx, sampling_mode='full')
        assert result['success'] is True
        assert result['statistics']['displayed_rows'] == 250
        assert result['sampling_info']['is_truncated'] is False

    def test_file_not_found(self):
        result = convert_xlsx_to_markdown('/nonexistent/file.xlsx')
        assert result['success'] is False


# ---------------------------------------------------------------------------
# Tests: Defragmentasi XML
# ---------------------------------------------------------------------------

class TestDefragmentDocxXml:
    def test_defragment_normal_docx(self, sample_docx):
        """File normal seharusnya tetap valid setelah defragmentasi."""
        result = defragment_docx_xml(sample_docx)
        assert result['success'] is True
        assert 'fragments_merged' in result
        assert 'paragraphs_cleaned' in result

    def test_file_not_found(self):
        result = defragment_docx_xml('/nonexistent/file.docx')
        assert result['success'] is False

    def test_wrong_extension(self, temp_dir):
        path = os.path.join(temp_dir, 'test.xlsx')
        with open(path, 'w') as f:
            f.write('hello')
        result = defragment_docx_xml(path)
        assert result['success'] is False
        assert '.docx' in result['error']

    def test_docx_still_valid_after_defragment(self, sample_docx):
        """DOCX harus masih bisa dibaca oleh python-docx setelah defragmentasi."""
        defragment_docx_xml(sample_docx)

        # Harus tetap bisa dibaca
        from docx import Document
        doc = Document(sample_docx)
        assert len(doc.paragraphs) > 0


if __name__ == '__main__':
    pytest.main([__file__, '-v'])
