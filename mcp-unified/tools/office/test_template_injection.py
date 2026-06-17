"""
Unit Tests — Template Injection (Fase 2)

Menguji render template docx dengan docxtpl, validasi path traversal,
dan eksekusi Skill TemplateInjectionSkill dengan mocking basis data.
"""
import os
import sys
import tempfile
import shutil
import pytest
from unittest.mock import AsyncMock, patch, MagicMock
from pathlib import Path

# Pastikan mcp-unified bisa diimport
sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from tools.office.template_tools import validate_path, render_docx_template
from skills.office.template_injection_skill import TemplateInjectionSkill, format_tgl
from core.task import Task


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def temp_dir():
    """Temporary directory di bawah workspace."""
    # Selalu gunakan subfolder di bawah /home/aseps/MCP demi keamanan path validation
    base_dir = "/home/aseps/MCP/scratch/test_templates"
    os.makedirs(base_dir, exist_ok=True)
    d = tempfile.mkdtemp(dir=base_dir)
    yield d
    shutil.rmtree(d, ignore_errors=True)


@pytest.fixture
def sample_template(temp_dir):
    """Buat file template .docx sederhana untuk testing."""
    from docx import Document
    
    path = os.path.join(temp_dir, 'template.docx')
    doc = Document()
    doc.add_heading('Disposisi Surat', level=1)
    
    # Tambahkan paragraf dengan placeholder docxtpl (Jinja2)
    doc.add_paragraph('Dari: {{ direktorat }}')
    doc.add_paragraph('Nomor ND: {{ nomor_nd }}')
    doc.add_paragraph('Tanggal: {{ tanggal_surat }}')
    doc.add_paragraph('Perihal: {{ hal }}')
    doc.add_paragraph('No Agenda Ses: {{ no_agenda_ses }}')
    
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# Tests: Path Validation (Security)
# ---------------------------------------------------------------------------

class TestPathValidation:
    def test_allow_valid_workspace_path(self):
        valid_path = "/home/aseps/MCP/korespondensi-server/templates_doc/template_disposisi.docx"
        p = validate_path(valid_path)
        assert p.is_absolute()
        assert str(p) == valid_path
        
    def test_allow_tmp_path(self):
        valid_path = "/tmp/Disposisi_test.docx"
        p = validate_path(valid_path)
        assert p.is_absolute()
        assert str(p) == valid_path
        
    def test_reject_outside_prefix(self):
        invalid_path = "/etc/passwd"
        with pytest.raises(PermissionError) as excinfo:
            validate_path(invalid_path)
        assert "Access to path '/etc/passwd' is denied" in str(excinfo.value)
        
    def test_reject_path_traversal(self):
        invalid_path = "/home/aseps/MCP/../../etc/passwd"
        with pytest.raises(PermissionError) as excinfo:
            validate_path(invalid_path)
        assert "Access to path" in str(excinfo.value)
        
    def test_custom_allowed_prefixes(self):
        path = "/var/log/app.log"
        # Harus raise error karena tidak diizinkan secara default
        with pytest.raises(PermissionError):
            validate_path(path)
            
        # Harus sukses jika kita berikan prefix log
        p = validate_path(path, allowed_prefixes=["/var/log"])
        assert str(p) == path


# ---------------------------------------------------------------------------
# Tests: Template Rendering Tool
# ---------------------------------------------------------------------------

class TestRenderDocxTemplate:
    def test_successful_rendering(self, sample_template, temp_dir):
        output_path = os.path.join(temp_dir, 'output_rendered.docx')
        context = {
            'direktorat': 'Direktorat PUU',
            'nomor_nd': 'ND-1234/PUU/2026',
            'tanggal_surat': '25 Mei 2026',
            'hal': 'Rapat Harmonisasi RUU',
            'no_agenda_ses': '089-I'
        }
        
        result = render_docx_template(
            template_path=sample_template,
            output_path=output_path,
            context=context
        )
        
        assert result['success'] is True
        assert os.path.exists(output_path)
        assert result['output_path'] == output_path
        
        # Baca kembali dan cek apakah placeholder digantikan
        from docx import Document
        doc = Document(output_path)
        all_text = "\n".join([p.text for p in doc.paragraphs])
        assert 'Dari: Direktorat PUU' in all_text
        assert 'Nomor ND: ND-1234/PUU/2026' in all_text
        assert 'Tanggal: 25 Mei 2026' in all_text
        assert 'Perihal: Rapat Harmonisasi RUU' in all_text
        assert 'No Agenda Ses: 089-I' in all_text

    def test_template_not_found(self, temp_dir):
        output_path = os.path.join(temp_dir, 'output.docx')
        result = render_docx_template(
            template_path='/home/aseps/MCP/nonexistent_template.docx',
            output_path=output_path,
            context={}
        )
        assert result['success'] is False
        assert 'file not found' in result['error'].lower()


# ---------------------------------------------------------------------------
# Tests: Date Formatting
# ---------------------------------------------------------------------------

class TestDateFormatting:
    def test_format_date_object(self):
        from datetime import date
        d = date(2026, 5, 25)
        assert format_tgl(d) == "25 Mei 2026"
        
    def test_format_datetime_object(self):
        from datetime import datetime
        d = datetime(2026, 12, 1, 14, 30)
        assert format_tgl(d) == "01 Desember 2026"
        
    def test_format_iso_string(self):
        assert format_tgl("2026-04-09T12:00:00Z") == "09 April 2026"
        assert format_tgl("2026-08-17") == "17 Agustus 2026"
        
    def test_format_invalid_string(self):
        assert format_tgl("not-a-date") == "not-a-date"
        assert format_tgl(None) == "-"


# ---------------------------------------------------------------------------
# Tests: Skill (Database Mocking & Execution)
# ---------------------------------------------------------------------------

class TestTemplateInjectionSkill:
    @pytest.mark.asyncio
    @patch('asyncpg.connect')
    async def test_inject_from_db_success(self, mock_connect, sample_template, temp_dir):
        # Setup mock connection and database response
        mock_conn = AsyncMock()
        mock_connect.return_value = mock_conn
        
        # Gunakan kelas dict langsung agar dict(mock_row) bekerja dengan benar
        db_record = {
            'id': 42,
            'unique_id': '0327_000.4.2_834_bu_set',
            'dari': 'Sekretariat Utama',
            'nomor_nd': 'ND-834/BU/SET/2026',
            'tanggal_surat': '2026-05-12',
            'hal': 'Undangan Sosialisasi',
            'no_agenda_dispo': '112-I',
            'tanggal_diterima_puu': '2026-05-14',
            'agenda_puu': '112-I'
        }
        
        # asyncpg.Record mendukung pemetaan/dict conversion
        mock_conn.fetchrow.return_value = db_record
        
        # Inisialisasi Skill
        skill = TemplateInjectionSkill()
        output_path = os.path.join(temp_dir, 'output_db_test.docx')
        
        result = await skill.inject_template_from_db(
            unique_id='0327_000.4.2_834_bu_set',
            template_path=sample_template,
            output_path=output_path
        )
        
        assert result['success'] is True
        assert os.path.exists(output_path)
        assert result['context_used']['direktorat'] == 'Sekretariat Utama'
        assert result['context_used']['nomor_nd'] == 'ND-834/BU/SET/2026'
        assert result['context_used']['tanggal_surat'] == '12 Mei 2026'
        assert result['context_used']['tgl_diterima'] == '14 Mei 2026'
        
        # Verify DB connection was called
        mock_connect.assert_called_once()
        mock_conn.fetchrow.assert_called_once_with(
            "SELECT * FROM surat_masuk_puu_internal WHERE unique_id = $1",
            '0327_000.4.2_834_bu_set'
        )

    @pytest.mark.asyncio
    @patch('asyncpg.connect')
    async def test_inject_from_db_not_found(self, mock_connect, sample_template):
        mock_conn = AsyncMock()
        mock_connect.return_value = mock_conn
        mock_conn.fetchrow.return_value = None
        
        skill = TemplateInjectionSkill()
        result = await skill.inject_template_from_db(
            unique_id='nonexistent_id',
            template_path=sample_template
        )
        
        assert result['success'] is False
        assert "tidak ditemukan" in result['error']

    @pytest.mark.asyncio
    @patch('asyncpg.connect')
    async def test_skill_execution_via_task(self, mock_connect, sample_template, temp_dir):
        # Setup mock db
        mock_conn = AsyncMock()
        mock_connect.return_value = mock_conn
        
        db_record = {
            'id': 1,
            'unique_id': 'test_id',
            'dari': 'Humas',
            'nomor_nd': 'ND-1',
            'tanggal_surat': '2026-05-01',
            'hal': 'Test',
            'no_agenda_dispo': '1-I',
            'tanggal_diterima_puu': None,
            'agenda_puu': None
        }
        mock_conn.fetchrow.return_value = db_record
        
        output_path = os.path.join(temp_dir, 'output_task.docx')
        
        # Create Task
        task = Task(
            id="task-001",
            type="template_injection",
            payload={
                "action": "inject_from_db",
                "unique_id": "test_id",
                "template_path": sample_template,
                "output_path": output_path
            }
        )
        
        skill = TemplateInjectionSkill()
        task_result = await skill.execute(task)
        
        assert task_result.success is True
        assert task_result.data['success'] is True
        assert task_result.data['output_path'] == output_path
        
        # Cek tanggal diterima default layout space placeholder
        assert task_result.data['context_used']['tgl_diterima'] == "                     /               / 2026"
